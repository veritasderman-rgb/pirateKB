"""Stage 1: extract text from every document in D:/Data/Mrak into staging/ (one .txt + meta)."""
import hashlib, html, json, os, re, sys, zipfile, datetime as dt
from concurrent.futures import ProcessPoolExecutor, as_completed

HERE = os.path.dirname(os.path.abspath(__file__))
STAGE = os.path.join(HERE, "staging")
MRAK = "D:/Data/Mrak"
MAX_ROWS = 400          # řádků na list tabulky
MAX_PDF_PAGES = 400     # dál už jde o knihy, text stačí z úvodu


def md_table(rows):
    rows = [[("" if c is None else str(c)).replace("\n", " ").replace("|", "\\|").strip() for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    keep = [i for i in range(width) if any(r[i] for r in rows)]
    rows = [[r[i] for i in keep] for r in rows]
    if not keep:
        return ""
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(keep)]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out)


def x_pdf(p):
    import pymupdf
    d = pymupdf.open(p)
    n = d.page_count
    parts = []
    for i, page in enumerate(d):
        if i >= MAX_PDF_PAGES:
            parts.append(f"\n[… text zkrácen, PDF má {n} stran …]")
            break
        parts.append(page.get_text("text"))
    meta = d.metadata or {}
    return "\n\n".join(parts), {"stran": n, "pdf_title": meta.get("title") or ""}


def x_docx(p):
    import docx
    d = docx.Document(p)
    out = []
    body = d.element.body
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for el in body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            par = Paragraph(el, d)
            t = par.text.strip()
            if not t:
                continue
            st = (par.style.name or "").lower() if par.style is not None else ""
            m = re.match(r"(heading|nadpis)\s*(\d)", st)
            if m:
                out.append("#" * min(int(m.group(2)) + 1, 6) + " " + t)
            elif "title" in st or "název" in st:
                out.append("## " + t)
            elif "list" in st:
                out.append("- " + t)
            else:
                out.append(t)
        elif tag == "tbl":
            tbl = Table(el, d)
            out.append(md_table([[c.text for c in r.cells] for r in tbl.rows]))
    return "\n\n".join(out), {}


ODF_NS = {
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0",
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
}


def _odf_text(el):
    """Plain text of an ODF element (spaces, tabs, line breaks)."""
    T = "{%s}" % ODF_NS["text"]
    parts = [el.text or ""]
    for ch in el:
        tag = ch.tag
        if tag == T + "s":
            parts.append(" " * int(ch.get(T + "c", "1")))
        elif tag == T + "tab":
            parts.append("\t")
        elif tag == T + "line-break":
            parts.append("\n")
        elif tag in (T + "note", ):
            pass
        else:
            parts.append(_odf_text(ch))
        parts.append(ch.tail or "")
    return "".join(parts)


def x_odf(p):
    import xml.etree.ElementTree as ET
    with zipfile.ZipFile(p) as z:
        root = ET.fromstring(z.read("content.xml"))
    T, TB, D = ("{%s}" % ODF_NS[k] for k in ("text", "table", "draw"))
    out = []
    ext = os.path.splitext(p)[1].lower()
    body = root.find("{%s}body" % ODF_NS["office"])

    def walk(el, depth=0):
        for ch in el:
            tag = ch.tag
            if tag == T + "h":
                lvl = int(ch.get(T + "outline-level", "1"))
                t = _odf_text(ch).strip()
                if t:
                    out.append("#" * min(lvl + 1, 6) + " " + t)
            elif tag == T + "p":
                t = _odf_text(ch).strip()
                if t:
                    out.append(t)
            elif tag == T + "list":
                for it in ch.iter(T + "list-item"):
                    for pp in it:
                        if pp.tag in (T + "p", T + "h"):
                            t = _odf_text(pp).strip()
                            if t:
                                out.append("- " + t)
            elif tag == TB + "table":
                rows = []
                for r in ch.iter(TB + "table-row"):
                    rep_r = int(r.get(TB + "number-rows-repeated", "1"))
                    cells = []
                    for c in r:
                        if c.tag not in (TB + "table-cell", TB + "covered-table-cell"):
                            continue
                        rep = int(c.get(TB + "number-columns-repeated", "1"))
                        txt = " ".join(_odf_text(pp).strip() for pp in c.iter(T + "p"))
                        cells += [txt] * min(rep, 50)
                    if any(cells):
                        rows += [cells] * min(rep_r, 3)
                    if len(rows) >= MAX_ROWS:
                        break
                name = ch.get(TB + "name", "")
                tab = md_table(rows)
                if tab:
                    if ext == ".ods":
                        out.append(f"## List: {name}")
                    out.append(tab)
            elif tag == D + "page":
                out.append(f"\n## Snímek {ch.get('{%s}name' % ODF_NS['draw'], '')}")
                walk(ch, depth + 1)
            else:
                walk(ch, depth + 1)
    walk(body)
    return "\n\n".join(out), {}


def x_xlsx(p):
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            if any(v is not None and str(v).strip() for v in r):
                rows.append(list(r)[:60])
            if len(rows) >= MAX_ROWS:
                break
        tab = md_table(rows)
        if tab:
            out.append(f"## List: {ws.title}\n\n{tab}")
    return "\n\n".join(out), {}


def x_pptx(p):
    import pptx
    pr = pptx.Presentation(p)
    out = []
    for i, s in enumerate(pr.slides, 1):
        texts = []
        for sh in s.shapes:
            if sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t:
                    texts.append(t)
            elif getattr(sh, "has_table", False) and sh.has_table:
                texts.append(md_table([[c.text for c in r.cells] for r in sh.table.rows]))
        if texts:
            out.append(f"## Snímek {i}\n\n" + "\n\n".join(texts))
    return "\n\n".join(out), {"snimku": len(pr.slides)}


def x_plain(p):
    raw = open(p, "rb").read()
    for enc in ("utf-8-sig", "cp1250", "latin-1"):
        try:
            return raw.decode(enc), {}
        except UnicodeDecodeError:
            pass
    return raw.decode("utf-8", "replace"), {}


def x_html(p):
    t, _ = x_plain(p)
    if t.count("<script") > 3 and len(re.sub(r"<script.*?</script>", "", t, flags=re.S)) < len(t) * 0.2:
        return "", {"pozn": "html je převážně skript (vizualizace)"}
    t = re.sub(r"<(script|style).*?</\1>", " ", t, flags=re.S | re.I)
    t = re.sub(r"<br\s*/?>|</p>|</h\d>|</li>|</tr>", "\n", t, flags=re.I)
    t = re.sub(r"<[^>]+>", " ", t)
    return html.unescape(t), {}


def x_rtf(p):
    t, _ = x_plain(p)
    t = re.sub(r"\\'([0-9a-f]{2})", lambda m: bytes([int(m.group(1), 16)]).decode("cp1250", "replace"), t)
    t = re.sub(r"\\u(-?\d+)\??", lambda m: chr(int(m.group(1)) % 65536), t)
    t = re.sub(r"\{\\\*[^{}]*\}", "", t)
    t = re.sub(r"\\par[d]?", "\n", t)
    t = re.sub(r"\\[a-z]+-?\d* ?", "", t)
    return t.replace("{", "").replace("}", ""), {}


def x_vtt(p):
    t, _ = x_plain(p)
    lines = [l for l in t.splitlines() if l.strip() and "-->" not in l and not l.strip().isdigit()
             and l.strip() != "WEBVTT"]
    return "\n".join(lines), {}


EXTRACT = {".pdf": x_pdf, ".docx": x_docx, ".odt": x_odf, ".ods": x_odf, ".odp": x_odf, ".odg": x_odf,
           ".xlsx": x_xlsx, ".pptx": x_pptx, ".md": x_plain, ".txt": x_plain, ".csv": x_plain,
           ".html": x_html, ".rtf": x_rtf, ".vtt": x_vtt, ".srt": x_vtt}


def work(item):
    path, ext, size = item
    full = os.path.join(MRAK, path)
    h = hashlib.sha256(open(full, "rb").read()).hexdigest()
    rec = {"path": path, "ext": ext, "size": size, "sha256": h,
           "mtime": dt.datetime.fromtimestamp(os.path.getmtime(full)).isoformat(timespec="seconds")}
    fn = EXTRACT.get(ext)
    if fn is None:
        rec["status"] = "nepodporovany-format"
    else:
        try:
            text, extra = fn(full)
            rec.update(extra)
            text = re.sub(r"[ \t]+\n", "\n", text)
            text = re.sub(r"\n{3,}", "\n\n", text).strip()
            rec["chars"] = len(text)
            rec["status"] = "ok" if len(text) >= 80 else "prazdne"
            if ext == ".pdf" and rec.get("stran") and len(text) / rec["stran"] < 40:
                rec["status"] = "sken"
            open(os.path.join(STAGE, h + ".txt"), "w", encoding="utf-8").write(text)
        except Exception as e:  # noqa
            rec["status"] = "chyba"
            rec["chyba"] = f"{type(e).__name__}: {e}"[:300]
    return rec


if __name__ == "__main__":
    os.makedirs(STAGE, exist_ok=True)
    inv = json.load(open(os.path.join(HERE, "inv.json"), encoding="utf-8"))
    # Mrak export obsahuje dvojitou složku (Assets/Assets/...): cesty jsou relativní k MRAK
    out = []
    with ProcessPoolExecutor(max_workers=8) as ex:
        futs = [ex.submit(work, it) for it in inv]
        for i, f in enumerate(as_completed(futs)):
            out.append(f.result())
            if i % 250 == 0:
                print(i, file=sys.stderr, flush=True)
    json.dump(out, open(os.path.join(HERE, "extracted.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=0)
    from collections import Counter
    print(Counter(r["status"] for r in out))
    print(Counter((r["ext"], r["status"]) for r in out if r["status"] != "ok"))
